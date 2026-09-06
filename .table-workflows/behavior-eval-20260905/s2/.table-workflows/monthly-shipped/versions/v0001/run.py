import argparse
import hashlib
import importlib.metadata
import json
import math
import platform
import subprocess
import sys
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
VERSION = ROOT / 'versions/v0001'


def now():
    return datetime.now(timezone.utc).isoformat()


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def save(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def structure(path):
    from openpyxl import load_workbook
    from xml.etree import ElementTree as ET
    def xml_value(value):
        return ET.tostring(value.to_tree(),encoding='unicode')
    workbook = load_workbook(path, data_only=False)
    with zipfile.ZipFile(path) as archive:
        # 记录全部对象与关系，避免库丢弃未支持对象后误判结构一致。
        parts = sorted(n for n in archive.namelist() if not n.startswith('docProps/') and n != 'xl/sharedStrings.xml')
        relationships = {}
        for name in parts:
            if name.endswith('.rels'):
                relationships[name] = sorted([dict(e.attrib) for e in ET.fromstring(archive.read(name))], key=lambda e:e.get('Id',''))
        feature_parts = {n:hashlib.sha256(archive.read(n)).hexdigest() for n in parts if any(x in n for x in ('drawings/','charts/','media/','externalLinks/','pivot','embeddings/','vba'))}
    sheets = []
    for sheet in workbook:
        sheets.append({
            'name':sheet.title, 'state':sheet.sheet_state, 'headers':[c.value for c in sheet[1]],
            'columnCount':sheet.max_column, 'merges':sorted(str(r) for r in sheet.merged_cells.ranges),
            'hiddenRows':sorted(i for i,d in sheet.row_dimensions.items() if d.hidden),
            'hiddenColumns':sorted(i for i,d in sheet.column_dimensions.items() if d.hidden),
            'formulas':[(c.coordinate,c.value) for row in sheet for c in row if c.data_type=='f'],
            'tables':[xml_value(t) for t in sheet.tables.values()],
            'validations':xml_value(sheet.data_validations),
            'conditionalFormats':[{'range':str(k.sqref),'rules':[xml_value(r) for r in sheet.conditional_formatting[k]]} for k in sheet.conditional_formatting],
            'freezePanes':sheet.freeze_panes, 'autoFilter':xml_value(sheet.auto_filter),
            'protection':xml_value(sheet.protection),
        })
    return {'fileType':'xlsx', 'sheets':sheets, 'definedNames':[xml_value(v) for v in workbook.defined_names.values()],
            'parts':parts, 'relationships':relationships, 'featureParts':feature_parts,
            'allowedVariance':['ordinary data values','data row count','cell styles and dimensions','document properties'],
            'fieldTypes':{'订单号':'nonempty text unique','仓库代码':'nonempty text','件数':'finite number','发货状态':'nonempty text','收款状态':'ignored'}}


def verify_dependencies(lock):
    if platform.python_version() != lock['pythonVersion']:
        raise ValueError('Python 版本变化，需检查模板环境')
    for item in lock['pythonPackages']:
        if importlib.metadata.version(item['name']) != item['version']:
            raise ValueError('Python 依赖版本变化: '+item['name'])
    for item in lock['nodePackages']:
        actual=json.loads(Path(item['path'],'package.json').read_text())
        if actual['version'] != item['version']:
            raise ValueError('Node 依赖版本变化: '+item['name'])
    for path, expected in lock['files'].items():
        if sha(path) != expected:
            raise ValueError('执行依赖内容变化: '+path)
    if subprocess.check_output([lock['nodePath'],'--version'],text=True).strip()!=lock['nodeVersion']:
        raise ValueError('Node 版本变化')


def run(request_path):
    request_path=Path(request_path).resolve()
    receipt_path=request_path.with_suffix('.receipt.json')
    receipt={'workflowId':'monthly-shipped','workflowVersion':1,'startedAt':now(),'exitStatus':1,
             'observedStructures':[],'warnings':[],'exceptionCount':0,'manualChanges':[],'outputPath':None}
    try:
        request=json.loads(request_path.read_text())
        manifest=json.loads((ROOT/'manifest.json').read_text())
        snapshot=json.loads((VERSION/'version-manifest.json').read_text())
        for key in ('workflowId','version','implementation','inputSlots','target','parameters'):
            if manifest.get(key)!=snapshot.get(key):
                raise ValueError('当前模板与版本快照不一致')
        receipt['artifacts']=manifest['implementation']['artifacts']
        for artifact in receipt['artifacts']:
            if sha(ROOT/artifact['path']) != artifact['sha256']:
                raise ValueError('执行产物变化，需优化并验证新版本')
        lock=json.loads((VERSION/'dependencies.lock.json').read_text())
        verify_dependencies(lock)
        if request.get('workflowId')!='monthly-shipped' or request.get('workflowVersion')!=1 or request.get('schemaVersion')!='0.1':
            raise ValueError('运行请求版本不匹配')
        if request.get('parameters',{}) or 'targetTemplatePath' in request or set(request['inputs'])!={'orders'} or len(request['inputs']['orders'])!=1:
            raise ValueError('本版本只接受一张订单原表，无额外参数或结果模板')
        source=Path(request['inputs']['orders'][0])
        output=Path(request['outputPath'])
        if not source.is_absolute() or not output.is_absolute():
            raise ValueError('输入和输出需为绝对路径')
        if output.exists() or source.resolve()==output.resolve():
            raise ValueError('禁止覆盖输入或已有结果')
        receipt['inputHashes']={str(source):sha(source)}
        receipt['outputPath']=str(output)
        expected=manifest['inputSlots'][0]['structureFingerprint']['value']
        observed=structure(source)
        receipt['observedStructures'].append({'path':str(source),'fingerprint':digest(observed),'expected':expected,'match':digest(observed)==expected})
        if digest(observed)!=expected:
            raise ValueError('原表结构已变化，请进入优化模式确认；未生成结果')
        from openpyxl import load_workbook
        sheet=load_workbook(source,data_only=False).active
        totals=defaultdict(float)
        seen=set()
        excluded=0
        matched=0
        for rowno,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
            order,warehouse,count,shipped,paid=row
            if not isinstance(order,str) or not order.strip() or order in seen:
                raise ValueError(f'第 {rowno} 行订单号为空或重复，需确认处理口径')
            seen.add(order)
            if not isinstance(shipped,str) or not shipped.strip():
                raise ValueError(f'第 {rowno} 行发货状态为空，需确认')
            if shipped!='已发货':
                excluded+=1
                continue
            if not isinstance(warehouse,str) or not warehouse.strip():
                raise ValueError(f'第 {rowno} 行仓库代码为空，不能汇总')
            if isinstance(count,bool) or not isinstance(count,(int,float)) or not math.isfinite(count):
                raise ValueError(f'第 {rowno} 行件数不是有效数字')
            totals[warehouse]+=count
            matched+=1
        rows=[[warehouse,int(count) if count.is_integer() else count] for warehouse,count in sorted(totals.items())]
        payload={'rows':rows,'outputPath':str(output)}
        if request_path.stem=='historical':
            payload['previewPath']=str(request_path.with_suffix('.png'))
        payload_path=request_path.with_suffix('.render.json')
        save(payload_path,payload)
        output.parent.mkdir(parents=True,exist_ok=True)
        process=subprocess.run([lock['nodePath'],str(VERSION/'render.mjs'),str(payload_path)],capture_output=True,text=True)
        request_path.with_suffix('.render.log').write_text(process.stdout+process.stderr)
        if process.returncode:
            raise ValueError('表格生成失败，见本次运行日志')
        reopened=load_workbook(output,data_only=False)
        actual=list(reopened.active.values)
        if actual != [tuple(['仓库代码','发货件数'])]+[tuple(r) for r in rows]:
            raise ValueError('导出后内容验证失败')
        receipt.update({'exitStatus':0,'outputHash':sha(output),'inputRecords':len(seen),'includedRecords':matched,
                        'excludedRecords':excluded,'unmatchedRecords':0,'outputRecords':len(rows),'shippedTotal':sum(r[1] for r in rows),'reopened':True})
    except Exception as error:
        receipt['exceptionCount']=1
        receipt['blockingError']=str(error)
    receipt['finishedAt']=now()
    receipt['outputExists']=bool(receipt['outputPath'] and Path(receipt['outputPath']).exists())
    save(receipt_path,receipt)
    print(json.dumps(receipt,ensure_ascii=False))
    return receipt['exitStatus']


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--request',required=True)
    sys.exit(run(parser.parse_args().request))
