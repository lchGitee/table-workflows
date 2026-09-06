import hashlib
import importlib.util
import json
import subprocess
import sys
import zipfile
from collections import defaultdict
from copy import copy
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

RUNS = Path(__file__).resolve().parent / 'attempt-2'
ROOT = Path(__file__).resolve().parents[2]
WORKSPACE = ROOT.parents[1]
VERSION = ROOT / 'versions/v0002'
NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
SHEET_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
VERIFY_EXISTING = '--verify-existing' in sys.argv


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def variant(name, change):
    # 只从已经保存的合成样本生成测试变体，不复制用户原表。
    path = RUNS / f'{name}.xlsx'
    if VERIFY_EXISTING:
        assert path.exists()
        return path
    assert not path.exists()
    with zipfile.ZipFile(ROOT / 'runs/fixtures/data-variance.xlsx') as source:
        parts = {n:source.read(n) for n in source.namelist()}
    change(parts)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as output:
        for key, value in parts.items():
            output.writestr(key, value)
    return path


def rename_ids(parts):
    for name in list(parts):
        if not name.endswith('.rels'):
            continue
        tree = ET.fromstring(parts[name])
        mapping = {}
        for index, entry in enumerate(tree):
            mapping[entry.attrib['Id']] = f'new-local-id-{index}'
            entry.attrib['Id'] = mapping[entry.attrib['Id']]
        tree[:] = reversed(list(tree))
        parts[name] = ET.tostring(tree)
        if name != '_rels/.rels':
            owner = str(Path(name).parent.parent / Path(name).name[:-5])
            document = ET.fromstring(parts[owner])
            for element in document.iter():
                for attribute, value in list(element.attrib.items()):
                    if attribute.startswith('{' + NS + '}'):
                        element.attrib[attribute] = mapping[value]
            parts[owner] = ET.tostring(document)


def change_target(parts):
    name = 'xl/_rels/workbook.xml.rels'
    tree = ET.fromstring(parts[name])
    for entry in tree:
        if entry.attrib['Type'].endswith('/theme'):
            entry.attrib['Target'] = '/xl/styles.xml'
    parts[name] = ET.tostring(tree)


def dangling_binding(parts):
    tree = ET.fromstring(parts['xl/workbook.xml'])
    tree.find(f'.//{{{SHEET_NS}}}sheet').set('{' + NS + '}id', 'missing-relationship')
    parts['xl/workbook.xml'] = ET.tostring(tree)


def duplicate_order(parts):
    tree = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    first = tree.find(f'.//{{{SHEET_NS}}}c[@r="A2"]')
    second = tree.find(f'.//{{{SHEET_NS}}}c[@r="A3"]')
    import copy
    second.attrib.clear()
    second.attrib.update(first.attrib)
    second.set('r', 'A3')
    second[:] = [copy.deepcopy(element) for element in first]
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(tree)


def expected_rows(source):
    rows = list(load_workbook(source, data_only=False).active.values)
    totals = defaultdict(float)
    included = 0
    for order, warehouse, count, shipped, paid in rows[1:]:
        if shipped == '已发货':
            totals[warehouse] += count
            included += 1
    return [('仓库代码', '发货件数')] + sorted(totals.items()), len(rows)-1, included


def compare_gold(output, expected):
    actual = load_workbook(output, data_only=False)
    gold = load_workbook(expected, data_only=False)
    assert actual.sheetnames == gold.sheetnames
    comparison = []
    for actual_sheet, gold_sheet in zip(actual, gold):
        assert actual_sheet.max_row == gold_sheet.max_row
        assert actual_sheet.max_column == gold_sheet.max_column
        for row in gold_sheet:
            for cell in row:
                other = actual_sheet[cell.coordinate]
                item = {'sheet':gold_sheet.title, 'cell':cell.coordinate,
                        'valueEqual':cell.value == other.value,
                        'formulaEqual':cell.data_type == other.data_type,
                        'styleEqual':all(copy(getattr(cell, attr)) == copy(getattr(other, attr))
                                         for attr in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'))}
                comparison.append(item)
    assert all(c['valueEqual'] and c['formulaEqual'] and c['styleEqual'] for c in comparison)
    return comparison


def main():
    if not VERIFY_EXISTING:
        RUNS.mkdir(exist_ok=False)
    manifest = json.loads((VERSION / 'version-manifest.json').read_text())
    old_hashes = {str(p.relative_to(ROOT)):sha(p) for p in (ROOT / 'versions/v0001').rglob('*') if p.is_file()}
    originals = {str(p):sha(p) for p in (WORKSPACE/'inbox').glob('*.xlsx')}
    spec = importlib.util.spec_from_file_location('candidate', VERSION/'run.py')
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    runner.verify_dependencies(json.loads((VERSION/'dependencies.lock.json').read_text()))
    for artifact in manifest['implementation']['artifacts']:
        assert sha(ROOT/artifact['path']) == artifact['sha256']
    fixtures = [(name, variant(name, operation), False) for name, operation in
                [('relationship-target', change_target), ('dangling-binding', dangling_binding), ('duplicate-order', duplicate_order)]]
    fixtures += [('renamed-ids', variant('renamed-ids', rename_ids), True)]
    cases = [
        ('historical', WORKSPACE/'inbox/原表.xlsx', True),
        ('current-month', WORKSPACE/'inbox/本月原表.xlsx', True),
        ('data-variance', ROOT/'runs/fixtures/data-variance.xlsx', True),
        ('hidden-column', ROOT/'runs/fixtures/hidden-column.xlsx', False),
    ] + fixtures
    results = []
    for name, source, should_pass in cases:
        request = RUNS/f'{name}.json'
        output = RUNS/f'{name}-result.xlsx'
        if not VERIFY_EXISTING:
            assert not request.exists() and not output.exists()
            save(request, {'schemaVersion':'0.1','workflowId':'monthly-shipped','workflowVersion':2,
                           'inputs':{'orders':[str(source)]},'outputPath':str(output),'parameters':{}})
            command = [str(request) if value == '{request}' else value for value in manifest['implementation']['command']]
            result = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
        receipt = json.loads(request.with_suffix('.receipt.json').read_text())
        assert (receipt['exitStatus'] == 0) == should_pass, (name, receipt)
        if not VERIFY_EXISTING:
            assert receipt['exitStatus'] == result.returncode
        assert output.exists() == should_pass
        summary = {'case':name,'inputPath':str(source),'expected':'PASS' if should_pass else 'BLOCK before output',
                   'actualExitStatus':receipt['exitStatus'],'outputExists':output.exists(),
                   'observed':receipt['observedStructures'],'blockingError':receipt.get('blockingError')}
        if should_pass:
            expected, count, included = expected_rows(source)
            assert list(load_workbook(output, data_only=False).active.values) == expected
            assert receipt['inputRecords'] == count and receipt['includedRecords'] == included
            assert receipt['excludedRecords'] == count-included
            assert receipt['shippedTotal'] == sum(r[1] for r in expected[1:])
            assert receipt['exceptionCount'] == 0 and receipt['unmatchedRecords'] == 0
            assert receipt['outputStructure']['match']
            assert receipt['inputHashes'][str(source)] == sha(source)
            summary['invariants'] = receipt['checks']
            summary['counts'] = {k:receipt[k] for k in ('inputRecords','includedRecords','excludedRecords','outputRecords','shippedTotal','exceptionCount','unmatchedRecords')}
        results.append(summary)
        print(json.dumps(summary, ensure_ascii=False), flush=True)
    comparison = compare_gold(RUNS/'historical-result.xlsx', WORKSPACE/'inbox/我做好的.xlsx')
    for name, expected in old_hashes.items():
        assert sha(ROOT/name) == expected
    for path, expected in originals.items():
        assert sha(Path(path)) == expected
    report = {'result':'PASS','workflowId':'monthly-shipped','version':2,'status':'TRIAL',
              'historicalCases':1,'independentHistoricalCases':0,
              'userAuthorization':'修正任务模板，让以后同类文件也能直接运行；业务统计规则按原来，保留旧版本，处理本月原表。',
              'changeScope':'只规范化关系 Id，继续比较类型、目标和引用绑定；统计逻辑不变。',
              'structureTests':results,'historicalCellComparison':comparison,
              'valueDifferences':0,'formulaDifferences':0,'styleDifferences':0,
              'oldVersionUnchanged':True,'oldVersionHashes':old_hashes,'originalFilesUnchanged':True,
              'schemaValidation':{'fullMachineValidation':False,'requiredFieldsAndSemanticChecks':'PASS',
                                  'reason':'无完整 JSON Schema 校验器；检查必填字段、版本、路径、命令、案例、产物和依赖。'}}
    save(RUNS/'validation-report.json', report)
    print('VALIDATION PASS')


if __name__ == '__main__':
    main()
