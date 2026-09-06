import hashlib
import json
import subprocess
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

RUNS = Path(__file__).resolve().parent
ROOT = RUNS.parents[1]
WORKSPACE = ROOT.parents[1]
VERSION = ROOT / 'versions/v0003'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def main():
    manifest = json.loads((VERSION / 'version-manifest.json').read_text())
    assert json.loads((ROOT / 'manifest.json').read_text())['version'] == 2
    old_hashes = json.loads((RUNS / 'previous-version-hashes.json').read_text())
    original_paths = [WORKSPACE / 'inbox' / name for name in ('原表.xlsx', '我做好的.xlsx', '本月原表.xlsx')]
    original_paths += [ROOT / 'runs' / name for name in ('本月仓库发货汇总-v0002.xlsx', '本月仓库发货汇总-v0002-本次修正.xlsx')]
    originals = {str(p):sha(p) for p in original_paths}
    previous_fixtures = ROOT / 'runs/upgrade-v0002/attempt-2'
    # 预期值直接来自已确认历史结果和明确的测试算例，不复用候选筛选逻辑。
    cases = [
        ('historical', WORKSPACE / 'inbox/原表.xlsx', [('W01',2),('W02',3)], 4, 2),
        ('current-month', WORKSPACE / 'inbox/本月原表.xlsx', [('W01',12),('W02',1),('W03',7)], 5, 4),
        ('data-variance', ROOT / 'runs/fixtures/data-variance.xlsx', [('SYN-W1',13),('SYN-W2',11)], 3, 3),
        ('renamed-ids', previous_fixtures / 'renamed-ids.xlsx', [('SYN-W1',13),('SYN-W2',11)], 3, 3),
        ('policy-truth-table', RUNS / 'policy-truth-table.xlsx', [('SYN-W1',10),('SYN-W2',17)], 6, 4),
        ('hidden-column', ROOT / 'runs/fixtures/hidden-column.xlsx', None, None, None),
        ('relationship-target', previous_fixtures / 'relationship-target.xlsx', None, None, None),
        ('dangling-binding', previous_fixtures / 'dangling-binding.xlsx', None, None, None),
        ('duplicate-order', previous_fixtures / 'duplicate-order.xlsx', None, None, None),
    ]
    results = []
    for name, source, expected, input_count, included in cases:
        request = RUNS / f'{name}.json'
        output = RUNS / f'{name}-result.xlsx'
        assert not request.exists() and not output.exists()
        save(request, {'schemaVersion':'0.1', 'workflowId':'monthly-shipped', 'workflowVersion':3,
                       'inputs':{'orders':[str(source)]}, 'outputPath':str(output), 'parameters':{}})
        command = [str(request) if part == '{request}' else part for part in manifest['implementation']['command']]
        process = subprocess.run(command + ['--validate-candidate'], cwd=ROOT, capture_output=True, text=True)
        receipt = json.loads(request.with_suffix('.receipt.json').read_text())
        should_pass = expected is not None
        assert (process.returncode == 0) == should_pass, (name, receipt)
        assert receipt['exitStatus'] == process.returncode
        assert output.exists() == should_pass
        item = {'case':name, 'expected':'PASS' if should_pass else 'BLOCK_BEFORE_OUTPUT',
                'exitStatus':process.returncode, 'outputExists':output.exists(),
                'observedStructures':receipt['observedStructures'], 'blockingError':receipt.get('blockingError')}
        if should_pass:
            actual = list(load_workbook(output, data_only=False).active.values)
            assert actual == [('仓库代码','发货件数')] + expected, (name, actual)
            assert receipt['includedTotal'] == sum(r[1] for r in expected)
            assert receipt['inputRecords'] == input_count and receipt['includedRecords'] == included
            assert receipt['excludedRecords'] == input_count-included
            assert receipt['unmatchedRecords'] == 0 and receipt['exceptionCount'] == 0
            assert receipt['outputStructure']['match']
            assert receipt['inputHashes'][str(source)] == sha(source)
            assert receipt['checks']['shipment-or-paid-pending'] == 'PASS'
            item['expectedRows'] = expected
            item['counts'] = {k:receipt[k] for k in ('inputRecords','includedRecords','excludedRecords','includedTotal','unmatchedRecords','exceptionCount')}
        results.append(item)
        print(json.dumps(item, ensure_ascii=False), flush=True)
    gold = load_workbook(WORKSPACE/'inbox/我做好的.xlsx')
    actual = load_workbook(RUNS/'historical-result.xlsx')
    assert gold.sheetnames == actual.sheetnames
    comparison = []
    for gs, result in zip(gold, actual):
        assert gs.max_row == result.max_row and gs.max_column == result.max_column
        for row in gs:
            for cell in row:
                other = result[cell.coordinate]
                comparison.append({'sheet':gs.title, 'cell':cell.coordinate, 'valueEqual':cell.value == other.value,
                                   'formulaEqual':cell.data_type == other.data_type,
                                   'styleEqual':all(copy(getattr(cell,k)) == copy(getattr(other,k)) for k in ('font','fill','border','alignment','number_format','protection'))})
    assert all(c['valueEqual'] and c['formulaEqual'] and c['styleEqual'] for c in comparison)
    for file, expected in old_hashes.items():
        assert sha(ROOT/file) == expected
    for file, expected in originals.items():
        assert sha(file) == expected
    save(RUNS/'validation-report.json', {
        'result':'PASS', 'workflowId':'monthly-shipped', 'version':3, 'status':'TRIAL',
        'userAuthorization':'我想以后待发货但已经收款的订单也算进去，原来已发货的还是照样算，帮我把这个汇总规则改一下。',
        'effectiveScope':'从本次升级后的运行开始；保留已有历史结果。W01=6 的人工调整不进入长期规则。',
        'historicalCases':1, 'independentHistoricalCases':0, 'cases':results,
        'historicalCellComparison':comparison, 'valueDifferences':0, 'formulaDifferences':0, 'styleDifferences':0,
        'oldVersionsUnchanged':True, 'originalFilesUnchanged':True, 'originalHashes':originals,
        'schemaValidation':{'fullMachineValidation':False, 'reason':'运行环境无完整 JSON Schema 校验器，发布前逐项校验必填字段和语义。'},
    })
    print('VALIDATION PASS')


if __name__ == '__main__':
    main()
