import importlib.util
import json
from pathlib import Path
from openpyxl import load_workbook

BASE=Path(__file__).resolve().parent
ROOT=BASE/'.table-workflows/monthly-shipped'
VERSION=ROOT/'versions/v0001'
spec=importlib.util.spec_from_file_location('runner',VERSION/'run.py')
runner=importlib.util.module_from_spec(spec)
spec.loader.exec_module(runner)
expected_path=BASE/'inbox/我做好的.xlsx'
actual_path=ROOT/'runs/historical-result.xlsx'
expected=load_workbook(expected_path,data_only=False)
actual=load_workbook(actual_path,data_only=False)
assert expected.sheetnames==actual.sheetnames
cells=[]
for source_sheet,result_sheet in zip(expected,actual):
    assert (source_sheet.max_row,source_sheet.max_column)==(result_sheet.max_row,result_sheet.max_column)
    for row in source_sheet:
        for cell in row:
            result=result_sheet[cell.coordinate]
            style=lambda c:(c.font.name,c.font.sz,c.font.bold,c.font.color.rgb,c.number_format,c.fill.patternType,c.alignment.horizontal)
            item={'sheet':source_sheet.title,'cell':cell.coordinate,'valueEqual':cell.value==result.value,
                  'formulaEqual':(cell.data_type=='f',cell.value if cell.data_type=='f' else None)==(result.data_type=='f',result.value if result.data_type=='f' else None),
                  'visibleStyleEqual':style(cell)==style(result)}
            assert item['valueEqual'] and item['formulaEqual'] and item['visibleStyleEqual'],item
            cells.append(item)
    assert source_sheet.sheet_state==result_sheet.sheet_state
    assert list(source_sheet.merged_cells.ranges)==list(result_sheet.merged_cells.ranges)==[]
    assert not source_sheet._charts and not result_sheet._charts and not source_sheet._images and not result_sheet._images
    assert list(expected.defined_names)==list(actual.defined_names)==[]
    assert [source_sheet.column_dimensions[c].width for c in ('A','B')]==[result_sheet.column_dimensions[c].width for c in ('A','B')]
    assert [source_sheet.row_dimensions[i].height for i in range(1,4)]==[result_sheet.row_dimensions[i].height for i in range(1,4)]
receipts={name:json.loads((ROOT/'runs'/f'{name}.receipt.json').read_text()) for name in ('historical','data-variance','hidden-column')}
assert receipts['historical']['exitStatus']==0 and receipts['historical']['shippedTotal']==5
assert receipts['data-variance']['exitStatus']==0 and receipts['data-variance']['shippedTotal']==18
assert list(load_workbook(ROOT/'runs/data-variance-result.xlsx',data_only=True).active.values)==[('仓库代码','发货件数'),('SYN-W1',7),('SYN-W2',11)]
assert receipts['hidden-column']['exitStatus']==1 and receipts['hidden-column']['outputExists'] is False
assert receipts['historical']['inputHashes'][str(BASE/'inbox/原表.xlsx')]==runner.sha(BASE/'inbox/原表.xlsx')
manifest=json.loads((ROOT/'manifest.json').read_text())
schema=json.loads(Path('/Users/lch/Desktop/project/table-workflows/references/manifest.schema.json').read_text())
# 环境没有完整 JSON Schema 校验器，按协议检查必填字段和主要语义约束。
assert set(schema['required'])<=set(manifest)
assert set(manifest)<=set(schema['properties'])
for key,defn in [('inputSlots','inputSlot')]:
    for value in manifest[key]:
        assert set(schema['$defs'][defn]['required'])<=set(value)
for key in ('target','implementation','validation'):
    assert set(schema['$defs'][key]['required'])<=set(manifest[key])
assert manifest['implementation']['command'].count('{request}')==1
assert manifest['version']==1 and manifest['workflowId']==ROOT.name
for artifact in manifest['implementation']['artifacts']:
    assert runner.sha(ROOT/artifact['path'])==artifact['sha256']
    assert (ROOT/artifact['path']).resolve().is_relative_to(ROOT.resolve())
for slot in manifest['inputSlots']:
    assert runner.digest(json.loads((ROOT/slot['structureFingerprint']['basis']).read_text()))==slot['structureFingerprint']['value']
timestamp=runner.now()
manifest['status']='TRIAL'
manifest['updatedAt']=timestamp
manifest['validation']['lastValidatedAt']=timestamp
manifest['validation']['goldCases'][0].update({'result':'PASS','lastRunAt':timestamp,'acceptedDifferenceCount':0})
runner.save(ROOT/'manifest.json',manifest)
runner.save(VERSION/'version-manifest.json',manifest)
report={'result':'PASS','workflowId':'monthly-shipped','version':1,'status':'TRIAL','historicalCases':1,'independentHistoricalCases':0,
        'policyConfirmation':'用户确认：只看发货状态，已发货就算进去，没收款也算。收款状态不影响这个统计。',
        'comparison':{'cells':cells,'valueDifferences':0,'formulaDifferences':0,'visibleStyleDifferences':0,'unexpectedCells':0},
        'totals':{'sourceRecords':4,'includedRecords':2,'excludedRecords':2,'unmatchedRecords':0,'exceptions':0,'warehouses':2,'shippedPieces':5},
        'structureTests':[{'case':name,'inputPath':json.loads((ROOT/'runs'/f'{name}.json').read_text())['inputs']['orders'][0],
                           'expected':'BLOCK before output' if name=='hidden-column' else 'PASS',
                           'actualExitStatus':receipt['exitStatus'],'outputExists':receipt['outputExists'],
                           'observed':receipt['observedStructures']} for name,receipt in receipts.items()],
        'reopened':True,'visualReview':'PASS: A1:B3 内容清晰，无截断；Arial 11、行高 23、列宽 18 与样本一致。',
        'originalFilesCopied':False,'expectedResultHash':runner.sha(expected_path),
        'schemaValidation':{'fullMachineValidation':False,'requiredFieldsAndSemanticChecks':'PASS','reason':'本地环境没有 JSON Schema 校验器，已逐项检查必填字段及主要语义约束。'}}
runner.save(ROOT/'runs/validation-report.json',report)
print(json.dumps({'status':'TRIAL','cellsCompared':len(cells),'differences':0,'fullSchemaValidation':False,'structuralTests':'PASS'},ensure_ascii=False))
