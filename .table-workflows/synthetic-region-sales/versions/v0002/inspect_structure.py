#!/usr/bin/env python3
import argparse
import json
import math
import os
import platform
import zipfile

import openpyxl
from openpyxl import load_workbook


def require(condition, message):
    if not condition:
        raise ValueError(message)


def value_matches(value, expected_type):
    if value is None:
        return True
    if expected_type in {"text", "text-id"}:
        return isinstance(value, str)
    if expected_type == "integer":
        return isinstance(value, int) and not isinstance(value, bool)
    if expected_type in {"finite-number", "number", "currency-2dp"}:
        return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)
    return False


def drawing_inventory(file_path, workbook):
    drawings = []
    for sheet in workbook.worksheets:
        drawings.extend(f"{sheet.title}:chart:{index}" for index, _ in enumerate(sheet._charts, start=1))
        drawings.extend(f"{sheet.title}:image:{index}" for index, _ in enumerate(sheet._images, start=1))
    with zipfile.ZipFile(file_path) as archive:
        drawing_parts = sorted(
            name
            for name in archive.namelist()
            if name.startswith("xl/drawings/") and name.endswith(".xml")
        )
    drawings.extend(f"package:{name}" for name in drawing_parts)
    return sorted(drawings)


def named_range_inventory(workbook):
    return sorted(
        [
            {"name": name, "definition": definition.attr_text}
            for name, definition in workbook.defined_names.items()
        ],
        key=lambda item: (item["name"], item["definition"] or ""),
    )


def inspect_structure(file_path, contract_path):
    with open(contract_path, "r", encoding="utf-8") as handle:
        contract = json.load(handle)

    workbook = load_workbook(file_path, data_only=False, read_only=False, keep_links=True)
    expected_sheet = contract["worksheets"][0]["name"]
    require(expected_sheet in workbook.sheetnames, f"缺少工作表：{expected_sheet}")
    sheet = workbook[expected_sheet]
    headers = contract["dataRegion"]["headers"]
    expected_types = contract["dataRegion"]["columnTypes"]
    actual_headers = [sheet.cell(row=1, column=index).value for index in range(1, len(headers) + 1)]

    for column_index, expected_type in enumerate(expected_types, start=1):
        values = [sheet.cell(row=row, column=column_index).value for row in range(2, sheet.max_row + 1)]
        require(
            all(value_matches(value, expected_type) for value in values),
            f"{expected_sheet} 第 {column_index} 列类型不符合 {expected_type}",
        )

    observed = json.loads(json.dumps(contract, ensure_ascii=False))
    observed["fileType"] = os.path.splitext(file_path)[1].lower()
    observed["worksheets"] = [
        {"order": index, "name": item.title, "visibility": item.sheet_state}
        for index, item in enumerate(workbook.worksheets, start=1)
    ]
    observed["dataRegion"]["headers"] = actual_headers
    observed["formulas"] = sorted(
        [
            {"cell": f"{item.title}!{cell.coordinate}", "formula": cell.value}
            for item in workbook.worksheets
            for row in item.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ],
        key=lambda entry: entry["cell"],
    )
    observed["mergedCells"] = sorted(
        f"{item.title}!{merged}"
        for item in workbook.worksheets
        for merged in item.merged_cells.ranges
    )
    observed["hiddenRows"] = sorted(
        f"{item.title}!{index}"
        for item in workbook.worksheets
        for index, dimension in item.row_dimensions.items()
        if dimension.hidden
    )
    observed["hiddenColumns"] = sorted(
        f"{item.title}!{key}"
        for item in workbook.worksheets
        for key, dimension in item.column_dimensions.items()
        if dimension.hidden
    )
    observed["namedRanges"] = named_range_inventory(workbook)
    observed["drawings"] = drawing_inventory(file_path, workbook)
    if "protectedScopes" in observed:
        observed["protectedScopes"] = sorted(
            item.title for item in workbook.worksheets if item.protection.sheet
        )
    return observed


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--versions", action="store_true")
    parser.add_argument("--file")
    parser.add_argument("--contract")
    args = parser.parse_args()

    if args.versions:
        print(json.dumps({
            "python": platform.python_version(),
            "openpyxl": openpyxl.__version__,
        }, sort_keys=True))
        return

    require(args.file and args.contract, "必须同时提供 --file 和 --contract")
    print(json.dumps(
        inspect_structure(os.path.abspath(args.file), os.path.abspath(args.contract)),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ))


if __name__ == "__main__":
    main()
